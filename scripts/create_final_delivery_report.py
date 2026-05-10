from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "FINAL_DELIVERY_REPORT_EXPANDED.docx"


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, size=8):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None, font_size=8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=font_size)
        shade_cell(table.rows[0].cells[i], "D9EAF7")
        if widths:
            table.rows[0].cells[i].width = widths[i]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size)
            if widths:
                cells[i].width = widths[i]
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.color.rgb = RGBColor(31, 78, 121)


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.05


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

doc.styles["Normal"].font.name = "Arial"
doc.styles["Normal"].font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Final Project Delivery Report")
run.bold = True
run.font.size = Pt(17)
run.font.color.rgb = RGBColor(31, 78, 121)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Software Development Case Studies\n").bold = True
meta.add_run("Project: Personal Productivity App (Notion-style)\n")
meta.add_run("Team: Zhanbolat Baxen, Tastan Magzhan\n")
meta.add_run("Repository: https://github.com/baksenf1/notion_like\n")
meta.add_run("Video: TODO - insert unlisted YouTube link")

add_heading(doc, "1. Product Summary", 1)
add_body(doc, "Personal Productivity App (Notion-style) is a web application for organizing notes, pages, workspaces, projects, members, and tasks in one place. It addresses the problem that students and small teams often coordinate their work through several separate tools, such as chats, documents, spreadsheets, and calendars. The primary user is a student or small coursework team that needs a simple shared workspace for planning and tracking project work.")
add_body(doc, "The delivered product is not intended to be a full Notion replacement. Its final scope is a working MVP that demonstrates authentication, private workspace access, page editing, project creation, member visibility, task assignment, task status tracking, and saved data through MongoDB.")

add_heading(doc, "2. Requirements Delivery", 1)
add_table(
    doc,
    ["ID", "Requirement", "Type", "Final Status", "Evidence / Note", "Sprint"],
    [
        ["FR-01", "Create, edit, and delete notes", "Functional", "Implemented", "Implemented as workspace pages with title and content. The user can open a page, edit the title/content, save the page, and remove pages from the workspace. Evidence: demo timestamp TODO; commit c512b47.", "Sprint 2"],
        ["FR-02", "Create and manage tasks", "Functional", "Implemented", "Task creation, editing, deletion, assignment, priority, status, due date, and table display are implemented. The workflow uses real form input and persists the created task. Evidence: demo timestamp TODO; commit bba9e4f.", "Sprint 2"],
        ["FR-03", "Mark tasks as completed", "Functional", "Implemented", "Completion is represented by changing task status to DONE. This matches the original requirement because the task state changes and the table reflects the result. Evidence: demo timestamp TODO.", "Sprint 2"],
        ["FR-04", "Organize content into pages", "Functional", "Partial", "Pages exist as workspace content and are accessible through navigation. However, tasks are organized mainly by projects/workspaces rather than embedded inside pages, so the full original page/task organization idea is only partially met.", "Sprint 3"],
        ["FR-05", "Search notes and tasks", "Functional", "Partial", "Task keyword search/filtering is implemented through the task table and backend query filters. Full unified search across page content and task content together is not completed, so this requirement remains partial.", "Sprint 3"],
        ["FR-06", "Save user data", "Functional", "Implemented", "User, workspace, page, project, member, and task data are persisted in MongoDB. The demo should show creating data, refreshing/reopening the view, and seeing that the data remains available.", "Sprint 3"],
        ["NFR-01", "Pages load and save within 4-7 seconds", "Non-functional", "Partial", "The app runs locally with acceptable loading during the demo. Formal performance testing on other networks and with larger datasets was not completed, so it cannot be marked fully implemented.", "Sprint 4"],
        ["NFR-02", "Usability: simple and easy for new users", "Non-functional", "Partial", "The UI has sidebar navigation, forms, tables, dialogs, and clear page/task screens. However, no formal usability test report with multiple users was completed, so the status is partial.", "Sprint 4"],
        ["NFR-03", "Secure login and password storage", "Non-functional", "Implemented", "Passwords are hashed with bcrypt in the user model. Authentication uses Passport local/Google strategies and session-based access. Protected routes require authenticated users.", "Sprint 2"],
        ["NFR-04", "Reliability: save data and minimize data loss", "Non-functional", "Partial", "MongoDB persistence works for saved entities, but autosave and crash recovery were not fully implemented. Manual save for pages reduces accidental changes but does not fully prevent data loss.", "Sprint 3"],
        ["NFR-05", "Scalability for more notes, tasks, and users", "Non-functional", "Partial", "MongoDB, pagination, filtering, and query limits are used in several endpoints. However, no load test with large workspaces was completed, so scalability is a design direction rather than a validated result.", "Sprint 4"],
    ],
    [Inches(0.55), Inches(1.4), Inches(0.85), Inches(0.85), Inches(3.1), Inches(0.65)],
    font_size=8,
)
add_body(doc, "In total, 5 requirements were fully implemented, 6 were partially implemented, and 0 were fully descoped. The delivered product is a functional MVP for workspace, page, project, and task coordination, but it does not fully validate performance, scalability, autosave reliability, or unified search across all content. These unmet requirements do not destroy the MVP hypothesis, but they limit the product to a course-ready prototype rather than a production-ready productivity platform.")
add_body(doc, "The most important implemented requirements are FR-01, FR-02, FR-03, FR-06, and NFR-03 because they prove that the application can support authenticated users, persistent data, page editing, and task coordination. The partial requirements are mostly quality and expansion items. They matter for a production release, but they do not block the demonstration of the main MVP workflow.")

add_heading(doc, "3. MVP Hypothesis Outcome", 1)
add_table(
    doc,
    ["Question", "Answer"],
    [
        ["State the original hypothesis", "If we build a collaborative workspace where a student team can create a project, assign tasks with status, priority, and due date, and view recent progress in one dashboard, we believe 2-5 person student project teams will use it at least twice during a project week to coordinate work, because seeing ownership and deadlines in one shared place reduces coordination friction."],
        ["What evidence supports it?", "TODO: insert real testing evidence. Example: During a test with [classmate/course team], the user created a workspace, created a project, created a task, assigned the task, and changed task status without developer help. Observable behavior: they used the dashboard/task table to identify task ownership. Quote: '[insert exact quote]'. The demo also supports the hypothesis by showing real input and real output for the workspace/project/task workflow."],
        ["What did not validate?", "Long-term retention was not validated. The prototype proves that the core workflow can be completed, but it does not prove that the same team will return twice in a real project week. Unified search and autosave reliability were also not fully validated. Because of this, the MVP outcome should be described as technically validated but not fully market-validated."],
        ["What changed as a result?", "The focus moved from a broad Notion clone to a smaller workspace coordination tool focused on projects, members, pages, and task ownership. Real-time collaboration, advanced document blocks, reminders, file uploads, and full content search were left for later because they were not necessary for the first MVP test."],
    ],
    [Inches(1.7), Inches(5.6)],
    font_size=9,
)
add_body(doc, "The evidence standard for this section is important. The final report should not only say that the team believes the product is useful. Before submission, add at least one concrete testing observation, user quote, or usage metric. The strongest simple evidence would be a short test with a classmate who is not part of the team and a note of how long it took them to complete the workflow.")

add_heading(doc, "4. Development Process", 1)
add_table(
    doc,
    ["Item", "What to include"],
    [
        ["Methodology", "Scrum was selected in Assignment 2 because requirements were partially known, the timeline was fixed, and the team needed incremental delivery. In practice, the project followed lightweight Scrum: sprint goals were defined around setup, core feature build, small feature completion, testing, and final submission. The most practical part of Scrum for this team was prioritizing the backlog so that authentication, database schema, and CRUD workflows were finished before optional improvements."],
        ["Github", "Repository: https://github.com/baksenf1/notion_like. Commit history: https://github.com/baksenf1/notion_like/commits/main. Evidence commits: c512b47 for backend/API implementation and bba9e4f for client implementation/tracking fix. TODO: include screenshot of commit history with dates visible. TODO: open at least two commit diffs in the video or screenshots if required by the instructor."],
        ["Ceremonies/artefacts", "Sprint planning was used informally to decide priority order: database/auth first, notes/pages/tasks next, then testing and documentation. Sprint reviews were informal and based on whether the prototype could demonstrate a working user story. Formal retrospectives were not fully run because the team was small and the course timeline was short. TODO: include screenshot of sprint plan from Assignment 2 or a task board if available."],
        ["What you learned", "Scrum helped keep the project focused on core features instead of adding too many Notion-like extras. It also made the team recognize scope creep early: real-time collaboration, reminders, and advanced document blocks were postponed. The main thing to do differently would be committing smaller changes across more dates and keeping a visible sprint board, because final assessment requires evidence of distributed work."],
    ],
    [Inches(1.35), Inches(5.95)],
    font_size=9,
)
add_body(doc, "Important evidence note: the current local commit history visible in this repository is concentrated on 2026-05-10. If earlier commits exist in another repository, branch, or local copy, they should be linked here. If not, this is a weakness in the Development Process section because the final project instructions require work across multiple points in time.")
add_body(doc, "The development order still followed the dependency logic from Assignment 2. The database schema and backend models came before frontend workflows; authentication came before workspace data access; and task/page features came before final documentation. This order reduced rework because the React client could rely on stable API shapes once the backend was available.")

add_heading(doc, "5. Theory in Practice", 1)
add_table(
    doc,
    ["Item", "Required content"],
    [
        ["Concept", "Eric Ries, The Lean Startup (2011), Build-Measure-Learn / MVP loop. The specific concept applied is that an MVP is not a smaller version of every future feature; it is the smallest product that can test the riskiest assumption."],
        ["Where it appeared", "This appeared when deciding not to build a complete Notion clone. The original idea could have expanded into reminders, advanced note blocks, real-time editing, file uploads, and full search. Instead, the team prioritized the workflow that tested the core assumption: a user can create a workspace/project, create pages and tasks, assign ownership, and see progress."],
        ["Outcome", "Applying the concept helped reduce feature creep. The product still introduced some product debt by building both pages and project/task management, but the final product stayed focused enough to demonstrate a working MVP. The cost of not applying the concept more strictly is that unified search, autosave, and scalability testing remained partial rather than fully validated."],
    ],
    [Inches(1.35), Inches(5.95)],
    font_size=9,
)
add_body(doc, "This concept is visible in the final scope. The project has enough functionality to test the main workflow, but it intentionally does not claim to be a complete productivity ecosystem. That distinction matters because the course assessment asks whether the product tested the MVP hypothesis, not whether it copied every feature of an existing market leader.")

add_heading(doc, "6. Architecture Delta", 1)
add_table(
    doc,
    ["Component / Decision", "Original Design", "Final Implementation"],
    [
        ["Product scope", "Personal productivity app with notes, tasks, reminders/search, and organization features similar to Notion.", "Collaborative workspace app with authentication, workspaces, members, projects, pages, tasks, status/priority/due dates, and dashboard summaries. The scope shifted toward team coordination because the implemented project naturally supported shared workspaces and task ownership."],
        ["Notes model", "Notes as a core standalone feature.", "Implemented as Pages with title and content inside a workspace. This simplified the model and matched the Notion-style page concept while still satisfying the note creation/editing requirement."],
        ["Task organization", "Tasks planned as general personal tasks.", "Tasks are connected to workspaces and projects, with assignment to members. This better supports team project management and creates clearer evidence for the MVP hypothesis."],
        ["Authentication approach", "Assignment 2 mentioned JWT-based session auth and bcrypt.", "Final implementation uses Passport local/Google strategies, session handling, and bcrypt password hashing to support email/password and Google OAuth. This changed the technical design but preserved the security goal."],
        ["Search", "Search notes and tasks.", "Task filtering/search is implemented; full unified search across page content and tasks is not completed. This divergence happened because task coordination became the primary MVP workflow."],
        ["Reliability", "Autosave every 30 seconds and crash recovery were planned.", "Manual save for pages and MongoDB persistence are implemented; autosave/crash recovery were not completed due to time and MVP priority."],
    ],
    [Inches(1.55), Inches(2.35), Inches(3.4)],
    font_size=8,
)
add_body(doc, "These deltas show that the final architecture did not exactly match the first design. The most important change was the move from a personal note/task app to a collaborative workspace structure. This change strengthened the MVP demonstration because it created clearer user stories around workspaces, projects, members, and assigned tasks.")

doc.save(OUT)
print(OUT)
